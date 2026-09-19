import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, Optional
import pandas as pd
import numpy as np
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.app.core.config import REPORTS_DIR
from backend.app.services.profiling_service import generate_dataset_profile
from backend.app.services.ai_service import generate_heuristic_ai_insights
from backend.app.services.stats_service import compute_correlation_matrix
from backend.app.services.dataset_detector import detect_dataset_domain
from backend.app.services.analyzers import get_analyzer

def generate_pdf_report(
    df: pd.DataFrame,
    dataset_id: Optional[int] = 1,
    dataset_name: str = "Dataset",
    title: Optional[str] = None,
    domain_override: Optional[str] = None,
    column_mapping: Optional[Dict[str, Any]] = None,
    output_path: Optional[str] = None,
    domain: Optional[str] = None
) -> str:
    """Generate high-quality multi-page domain-adaptive executive PDF report with ReportLab."""
    effective_domain = domain_override or domain
    detected = detect_dataset_domain(df, dataset_name)
    if not effective_domain:
        effective_domain = detected["domain"]
        
    analyzer = get_analyzer(effective_domain)
    mapping = column_mapping or detected["detected_fields"]

    report_title = title or f"{analyzer.domain_display_name} Report — {dataset_name}"
    
    if output_path:
        file_path = Path(output_path)
        file_path.parent.mkdir(parents=True, exist_ok=True)
    else:
        unique_filename = f"report_{dataset_id or 1}_{effective_domain}_{uuid.uuid4().hex[:8]}.pdf"
        file_path = REPORTS_DIR / unique_filename
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    
    profile = generate_dataset_profile(df, dataset_id or 1, dataset_name)
    ai_insights = generate_heuristic_ai_insights(df, dataset_id or 1, dataset_name, domain_override=effective_domain, column_mapping=mapping)
    domain_data = analyzer.get_report_data(df, mapping)
    
    doc = SimpleDocTemplate(
        str(file_path),
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    PRIMARY = colors.HexColor("#1E293B")
    ACCENT = colors.HexColor("#4F46E5")
    SUCCESS = colors.HexColor("#059669")
    TEXT_DARK = colors.HexColor("#0F172A")
    MUTED = colors.HexColor("#64748B")
    BG_LIGHT = colors.HexColor("#F8FAFC")
    CARD_BG = colors.HexColor("#EEF2FF")
    
    title_style = ParagraphStyle("ReportTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=PRIMARY, spaceAfter=3)
    subtitle_style = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=MUTED, spaceAfter=12)
    h2_style = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=ACCENT, spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle("ReportBody", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_DARK)
    card_title_style = ParagraphStyle("CardTitle", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=ACCENT)
    
    story = []
    
    # Banner & Header
    story.append(Paragraph(f"AI DATASENSE — {analyzer.domain_display_name.upper()} REPORT", ParagraphStyle("Banner", fontName="Helvetica-Bold", fontSize=8, textColor=ACCENT, spaceAfter=3)))
    story.append(Paragraph(report_title, title_style))
    story.append(Paragraph(f"Generated on {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')} • Domain: <b>{analyzer.domain_display_name}</b> • Confidential", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=ACCENT, spaceAfter=12))
    
    # 1. Domain-Specific Executive Scorecard
    story.append(Paragraph("1. Executive Domain Scorecard", h2_style))
    overview_dict = domain_data.get("overview", {})
    overview_items = list(overview_dict.items())
    
    kpi_rows = []
    for i in range(0, len(overview_items), 3):
        row = []
        for k, v in overview_items[i:i+3]:
            row.append(Paragraph(f"<b>{k}:</b> {v}", body_style))
        while len(row) < 3:
            row.append(Paragraph("", body_style))
        kpi_rows.append(row)
        
    if not kpi_rows:
        kpi_rows = [
            [Paragraph(f"<b>Total Records:</b> {profile.rows:,}", body_style), Paragraph(f"<b>Columns:</b> {profile.columns}", body_style), Paragraph(f"<b>Quality Score:</b> {profile.data_quality_score}/100", body_style)]
        ]

    kpi_table = Table(kpi_rows, colWidths=[180, 180, 180])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), CARD_BG),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 8))
    
    # 2. Executive AI Synthesis
    story.append(Paragraph("2. Executive AI Synthesis", h2_style))
    story.append(Paragraph(ai_insights.executive_summary, body_style))
    story.append(Spacer(1, 8))
    
    # 3. Domain Specific Tables (Tailored for each domain)
    if effective_domain == "student" and "subject_table" in domain_data:
        story.append(Paragraph("3. Academic Subject Performance Analysis", h2_style))
        s_rows = [["Subject", "Average", "Highest", "Lowest", "Median", "Std Dev"]]
        for item in domain_data["subject_table"]:
            s_rows.append([
                str(item["Subject"]),
                f"{item['Average']:.1f}" if item['Average'] is not None else "-",
                f"{item['Highest']:.0f}" if item['Highest'] is not None else "-",
                f"{item['Lowest']:.0f}" if item['Lowest'] is not None else "-",
                f"{item['Median']:.1f}" if item['Median'] is not None else "-",
                f"{item['Std Dev']:.2f}" if item['Std Dev'] is not None else "-"
            ])
        s_table = Table(s_rows, colWidths=[140, 80, 80, 80, 80, 80])
        s_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(s_table)
        
        if domain_data.get("top_performers"):
            story.append(Spacer(1, 6))
            story.append(Paragraph("<b>Top Performing Students:</b>", card_title_style))
            for top_s in domain_data["top_performers"][:4]:
                story.append(Paragraph(f"• <b>{top_s['name']}</b> ({top_s.get('student_id', '')}): Average Score: <b>{top_s['average_marks']:.1f}</b> marks", body_style))

    elif effective_domain == "ecommerce" and domain_data.get("top_products"):
        story.append(Paragraph("3. Leading Products & Category Performance", h2_style))
        p_rows = [["Product", "Revenue (INR)", "Sales Share %"]]
        for p in domain_data["top_products"]:
            p_rows.append([str(p["product"]), f"₹{p['revenue']:,.2f}", f"{p['share']}%"])
        p_table = Table(p_rows, colWidths=[240, 150, 150])
        p_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(p_table)

    elif effective_domain == "hr" and domain_data.get("department_breakdown"):
        story.append(Paragraph("3. Departmental Workforce Breakdown", h2_style))
        d_rows = [["Department", "Headcount", "Headcount Share %", "Avg Salary (INR)"]]
        for d in domain_data["department_breakdown"]:
            d_rows.append([str(d["department"]), str(d["headcount"]), f"{d['percentage']}%", f"₹{d['avg_salary']:,.2f}"])
        d_table = Table(d_rows, colWidths=[180, 100, 120, 140])
        d_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(d_table)

    elif effective_domain == "banking" and domain_data.get("credit_tiers"):
        story.append(Paragraph("3. Credit Quality & Tier Distribution", h2_style))
        c_rows = [["Credit Tier", "Customer Count", "Portfolio Share %"]]
        for t in domain_data["credit_tiers"]:
            c_rows.append([str(t["tier"]), str(t["count"]), f"{t['percentage']}%"])
        c_table = Table(c_rows, colWidths=[220, 160, 160])
        c_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(c_table)

    elif effective_domain == "finance" and domain_data.get("entity_breakdown"):
        story.append(Paragraph("3. Financial Contribution by Business Entity / Category", h2_style))
        f_rows = [["Entity / Account", "Financial Volume (INR)", "Share %"]]
        for item in domain_data["entity_breakdown"]:
            f_rows.append([str(item["entity"]), f"₹{item['amount']:,.2f}", f"{item['share']}%"])
        f_table = Table(f_rows, colWidths=[240, 150, 150])
        f_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(f_table)

    elif effective_domain == "marketing" and domain_data.get("channel_breakdown"):
        story.append(Paragraph("3. Marketing Channel & Conversion Volume", h2_style))
        m_rows = [["Channel", "Traffic / Conversion Volume", "Share %"]]
        for item in domain_data["channel_breakdown"]:
            m_rows.append([str(item["channel"]), f"{item['volume']:,.0f}", f"{item['share']}%"])
        m_table = Table(m_rows, colWidths=[240, 150, 150])
        m_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(m_table)

    elif effective_domain == "healthcare" and domain_data.get("diagnosis_distribution"):
        story.append(Paragraph("3. Diagnosis & Clinical Condition Prevalence", h2_style))
        h_rows = [["Condition / Diagnosis", "Patient Count", "Prevalence Share %"]]
        for item in domain_data["diagnosis_distribution"]:
            h_rows.append([str(item["condition"]), str(item["count"]), f"{item['percentage']}%"])
        h_table = Table(h_rows, colWidths=[240, 150, 150])
        h_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(h_table)

    else:
        story.append(Paragraph("3. Descriptive Statistics Summary", h2_style))
        if profile.numeric_stats:
            table_rows = [["Feature", "Mean", "Std Dev", "Min", "Median (Q2)", "Max", "Skew"]]
            for ns in profile.numeric_stats[:6]:
                table_rows.append([
                    ns.name[:18],
                    f"{ns.mean:,.2f}" if ns.mean is not None else "-",
                    f"{ns.std:,.2f}" if ns.std is not None else "-",
                    f"{ns.min:,.2f}" if ns.min is not None else "-",
                    f"{ns.median:,.2f}" if ns.median is not None else "-",
                    f"{ns.max:,.2f}" if ns.max is not None else "-",
                    f"{ns.skewness:.2f}" if ns.skewness is not None else "-"
                ])
            stats_table = Table(table_rows, colWidths=[120, 70, 70, 70, 70, 70, 70])
            stats_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(stats_table)

    story.append(Spacer(1, 8))
    
    # 4. Strategic Recommendations
    story.append(Paragraph("4. AI Strategic Recommendations & Action Items", h2_style))
    for rec in ai_insights.strategic_recommendations:
        story.append(Paragraph(f"<b>[{rec.badge.upper()}] {rec.title}</b>", card_title_style))
        story.append(Paragraph(rec.summary, body_style))
        for bp in rec.bullet_points:
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;✓ {bp}", body_style))
        story.append(Spacer(1, 3))
        
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=0.5, color=MUTED))
    story.append(Paragraph("Generated automatically by AI DataSense — Universal AI Data Analytics Platform.", ParagraphStyle("Foot", fontName="Helvetica-Oblique", fontSize=7, textColor=MUTED, alignment=1)))

    doc.build(story)
    return str(file_path)

def generate_excel_report(
    df: pd.DataFrame,
    dataset_id: Optional[int] = 1,
    dataset_name: str = "Dataset",
    domain_override: Optional[str] = None,
    column_mapping: Optional[Dict[str, Any]] = None,
    output_path: Optional[str] = None,
    title: Optional[str] = None,
    domain: Optional[str] = None
) -> str:
    """Generate professional multi-sheet Excel workbook adapting to the dataset domain."""
    effective_domain = domain_override or domain
    detected = detect_dataset_domain(df, dataset_name)
    if not effective_domain:
        effective_domain = detected["domain"]
        
    analyzer = get_analyzer(effective_domain)
    mapping = column_mapping or detected["detected_fields"]

    if output_path:
        file_path = Path(output_path)
        file_path.parent.mkdir(parents=True, exist_ok=True)
    else:
        unique_filename = f"analytics_{dataset_id or 1}_{effective_domain}_{uuid.uuid4().hex[:8]}.xlsx"
        file_path = REPORTS_DIR / unique_filename
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    
    profile = generate_dataset_profile(df, dataset_id or 1, dataset_name)
    corr = compute_correlation_matrix(df)
    ai_insights = generate_heuristic_ai_insights(df, dataset_id or 1, dataset_name, domain_override=effective_domain, column_mapping=mapping)
    domain_data = analyzer.get_report_data(df, mapping)
    
    wb = openpyxl.Workbook()
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ---------------- SHEET 1: EXECUTIVE SUMMARY ----------------
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True
    
    sheet_title = title or f"AI DataSense — {analyzer.domain_display_name} Report: {dataset_name}"
    ws_summary["A1"] = sheet_title
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Domain: {analyzer.domain_display_name} | Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ws_summary["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")
    
    overview_dict = domain_data.get("overview", {})
    kpis = list(overview_dict.items()) if overview_dict else [
        ("Total Records", profile.rows),
        ("Total Columns", profile.columns),
        ("Data Quality Score", f"{profile.data_quality_score} / 100")
    ]
    
    ws_summary["A4"] = "Domain Metric"
    ws_summary["B4"] = "Value"
    ws_summary["A4"].fill = header_fill
    ws_summary["B4"].fill = header_fill
    ws_summary["A4"].font = header_font
    ws_summary["B4"].font = header_font
    
    for row_idx, (k, v) in enumerate(kpis, start=5):
        ws_summary[f"A{row_idx}"] = str(k)
        ws_summary[f"B{row_idx}"] = str(v)
        ws_summary[f"A{row_idx}"].border = border_thin
        ws_summary[f"B{row_idx}"].border = border_thin
        ws_summary[f"A{row_idx}"].font = bold_font
        ws_summary[f"B{row_idx}"].font = regular_font

    next_row = 5 + len(kpis) + 2
    ws_summary[f"A{next_row}"] = "AI Executive Synthesis:"
    ws_summary[f"A{next_row}"].font = bold_font
    ws_summary[f"A{next_row+1}"] = ai_insights.executive_summary
    ws_summary[f"A{next_row+1}"].font = regular_font

    # ---------------- SHEET 2: DOMAIN-SPECIFIC METRICS / STATS ----------------
    if effective_domain == "student" and "subject_table" in domain_data:
        ws_stats = wb.create_sheet(title="Subject Performance")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Subject", "Average", "Highest", "Lowest", "Median", "Std Dev"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for r_idx, item in enumerate(domain_data["subject_table"], start=2):
            row_vals = [item["Subject"], item["Average"], item["Highest"], item["Lowest"], item["Median"], item["Std Dev"]]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "ecommerce" and domain_data.get("top_products"):
        ws_stats = wb.create_sheet(title="Product Performance")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Product", "Revenue", "Sales Share %"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, p in enumerate(domain_data["top_products"], start=2):
            row_vals = [p["product"], p["revenue"], f"{p['share']}%"]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "hr" and domain_data.get("department_breakdown"):
        ws_stats = wb.create_sheet(title="Department Breakdown")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Department", "Headcount", "Percentage", "Avg Salary"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, d in enumerate(domain_data["department_breakdown"], start=2):
            row_vals = [d["department"], d["headcount"], f"{d['percentage']}%", d["avg_salary"]]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "banking" and domain_data.get("credit_tiers"):
        ws_stats = wb.create_sheet(title="Credit Tiers")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Credit Tier", "Customer Count", "Portfolio Share %"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, t in enumerate(domain_data["credit_tiers"], start=2):
            row_vals = [t["tier"], t["count"], f"{t['percentage']}%"]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "finance" and domain_data.get("entity_breakdown"):
        ws_stats = wb.create_sheet(title="Entity Breakdown")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Entity", "Amount", "Share %"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, item in enumerate(domain_data["entity_breakdown"], start=2):
            row_vals = [item["entity"], item["amount"], f"{item['share']}%"]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "marketing" and domain_data.get("channel_breakdown"):
        ws_stats = wb.create_sheet(title="Channel Breakdown")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Channel", "Volume", "Share %"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, item in enumerate(domain_data["channel_breakdown"], start=2):
            row_vals = [item["channel"], item["volume"], f"{item['share']}%"]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    elif effective_domain == "healthcare" and domain_data.get("diagnosis_distribution"):
        ws_stats = wb.create_sheet(title="Clinical Demographics")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Condition", "Count", "Percentage"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, item in enumerate(domain_data["diagnosis_distribution"], start=2):
            row_vals = [item["condition"], item["count"], f"{item['percentage']}%"]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    else:
        ws_stats = wb.create_sheet(title="Descriptive Statistics")
        ws_stats.views.sheetView[0].showGridLines = True
        headers = ["Feature", "Count", "Missing", "Mean", "Std Dev", "Min", "25%", "Median", "75%", "Max", "IQR", "Skewness"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_stats.cell(row=1, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for r_idx, ns in enumerate(profile.numeric_stats, start=2):
            row_vals = [ns.name, ns.count, ns.missing, ns.mean, ns.std, ns.min, ns.q25, ns.median, ns.q75, ns.max, ns.iqr, ns.skewness]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    # ---------------- SHEET 3: CORRELATION MATRIX ----------------
    ws_corr = wb.create_sheet(title="Correlation Matrix")
    ws_corr.views.sheetView[0].showGridLines = True
    if corr.columns and corr.matrix:
        ws_corr.cell(row=1, column=1, value="Feature").fill = header_fill
        ws_corr.cell(row=1, column=1).font = header_font
        for c_idx, col_name in enumerate(corr.columns, start=2):
            cell = ws_corr.cell(row=1, column=c_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
        for r_idx, (col_name, row_vals) in enumerate(zip(corr.columns, corr.matrix), start=2):
            ws_corr.cell(row=r_idx, column=1, value=col_name).font = bold_font
            ws_corr.cell(row=r_idx, column=1).border = border_thin
            for c_idx, val in enumerate(row_vals, start=2):
                cell = ws_corr.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                cell.font = regular_font

    # ---------------- SHEET 4: RAW DATASET ----------------
    ws_data = wb.create_sheet(title="Dataset Sample")
    ws_data.views.sheetView[0].showGridLines = True
    for c_idx, col_name in enumerate(df.columns, start=1):
        cell = ws_data.cell(row=1, column=c_idx, value=str(col_name))
        cell.fill = header_fill
        cell.font = header_font
    sample_df = df.head(1000)
    for r_idx, row in enumerate(sample_df.values, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_data.cell(row=r_idx, column=c_idx, value=str(val) if pd.notna(val) else "")
            cell.border = border_thin

    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(12, min(40, max_len + 3))
            
    wb.save(str(file_path))
    return str(file_path)
